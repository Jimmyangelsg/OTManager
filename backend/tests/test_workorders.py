"""Backend tests for Work Orders API."""
import os
import io
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get('REACT_APP_BACKEND_URL', 'https://orden-trabajo-local.preview.emergentagent.com').rstrip('/')
API = f"{BASE_URL}/api"


@pytest.fixture(scope="module")
def session():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    yield s
    # Cleanup TEST_ prefixed work orders
    try:
        resp = s.get(f"{API}/workorders?search=TEST_")
        if resp.status_code == 200:
            for wo in resp.json():
                if wo.get("ot_number", "").startswith("TEST_"):
                    s.delete(f"{API}/workorders/{wo['id']}")
    except Exception:
        pass


@pytest.fixture
def created_wo(session):
    """Create a single work order, yield, then delete."""
    payload = {
        "ot_number": "TEST_FIXT_001",
        "requestor": "Test User",
        "task_detail": "Fixture WO",
        "service_desk_number": "SD-FIXT",
        "observations": "fixture"
    }
    r = session.post(f"{API}/workorders", json=payload)
    assert r.status_code == 200, r.text
    wo = r.json()
    yield wo
    session.delete(f"{API}/workorders/{wo['id']}")


# --- CRUD ---
class TestCRUD:
    def test_root(self, session):
        r = session.get(f"{API}/")
        assert r.status_code == 200
        assert "message" in r.json()

    def test_create_returns_id_and_sort_order(self, session):
        payload = {"ot_number": "TEST_CREATE_001", "requestor": "Alice", "task_detail": "Create test"}
        r = session.post(f"{API}/workorders", json=payload)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "id" in data and len(data["id"]) > 0
        assert data["ot_number"] == "TEST_CREATE_001"
        assert data["requestor"] == "Alice"
        assert data["sort_order"] >= 1.0
        # cleanup
        session.delete(f"{API}/workorders/{data['id']}")

    def test_list_sorted_by_sort_order_desc(self, session):
        # Create two work orders
        a = session.post(f"{API}/workorders", json={"ot_number": "TEST_LIST_A"}).json()
        b = session.post(f"{API}/workorders", json={"ot_number": "TEST_LIST_B"}).json()
        try:
            r = session.get(f"{API}/workorders")
            assert r.status_code == 200
            all_wos = r.json()
            # b created after a -> b should appear before a
            idx_a = next(i for i, w in enumerate(all_wos) if w["id"] == a["id"])
            idx_b = next(i for i, w in enumerate(all_wos) if w["id"] == b["id"])
            assert idx_b < idx_a, "Most recently created OT should be at top"
            # also verify sort_order desc property generally
            orders = [w.get("sort_order", 0) for w in all_wos[:5]]
            assert orders == sorted(orders, reverse=True)
        finally:
            session.delete(f"{API}/workorders/{a['id']}")
            session.delete(f"{API}/workorders/{b['id']}")

    def test_get_single(self, session, created_wo):
        r = session.get(f"{API}/workorders/{created_wo['id']}")
        assert r.status_code == 200
        assert r.json()["ot_number"] == created_wo["ot_number"]

    def test_get_not_found(self, session):
        r = session.get(f"{API}/workorders/non-existent-id-xxx")
        assert r.status_code == 404

    def test_update_persists(self, session, created_wo):
        r = session.put(f"{API}/workorders/{created_wo['id']}", json={"requestor": "Updated Person"})
        assert r.status_code == 200
        assert r.json()["requestor"] == "Updated Person"
        # verify via GET
        g = session.get(f"{API}/workorders/{created_wo['id']}")
        assert g.json()["requestor"] == "Updated Person"
        # unchanged
        assert g.json()["ot_number"] == created_wo["ot_number"]

    def test_delete_removes(self, session):
        c = session.post(f"{API}/workorders", json={"ot_number": "TEST_DEL_001"}).json()
        r = session.delete(f"{API}/workorders/{c['id']}")
        assert r.status_code == 200
        g = session.get(f"{API}/workorders/{c['id']}")
        assert g.status_code == 404


# --- Search and filters ---
class TestSearchFilters:
    def test_search_by_ot_number(self, session):
        c = session.post(f"{API}/workorders", json={"ot_number": "TEST_SEARCH_XYZ", "requestor": "Bob"}).json()
        try:
            r = session.get(f"{API}/workorders?search=SEARCH_XYZ")
            assert r.status_code == 200
            results = r.json()
            assert any(w["id"] == c["id"] for w in results)
        finally:
            session.delete(f"{API}/workorders/{c['id']}")

    def test_filter_by_requestor(self, session):
        c = session.post(f"{API}/workorders", json={"ot_number": "TEST_REQ_001", "requestor": "UniqueRequestorXYZ"}).json()
        try:
            r = session.get(f"{API}/workorders?requestor=UniqueRequestorXYZ")
            assert r.status_code == 200
            results = r.json()
            assert len(results) >= 1
            assert any(w["id"] == c["id"] for w in results)
        finally:
            session.delete(f"{API}/workorders/{c['id']}")


# --- Reorder ---
class TestReorder:
    def test_reorder_assigns_sort_order(self, session):
        a = session.post(f"{API}/workorders", json={"ot_number": "TEST_REORD_A"}).json()
        b = session.post(f"{API}/workorders", json={"ot_number": "TEST_REORD_B"}).json()
        c = session.post(f"{API}/workorders", json={"ot_number": "TEST_REORD_C"}).json()
        try:
            # Reorder: a should be first (highest sort_order)
            ordered = [a["id"], b["id"], c["id"]]
            r = session.post(f"{API}/workorders/reorder", json={"ordered_ids": ordered})
            assert r.status_code == 200
            body = r.json()
            assert body["updated"] == 3
            # Verify via GET
            all_wos = session.get(f"{API}/workorders").json()
            ids_in_order = [w["id"] for w in all_wos if w["id"] in ordered]
            assert ids_in_order == ordered, f"Expected {ordered}, got {ids_in_order}"
        finally:
            for x in (a, b, c):
                session.delete(f"{API}/workorders/{x['id']}")

    def test_reorder_empty(self, session):
        r = session.post(f"{API}/workorders/reorder", json={"ordered_ids": []})
        assert r.status_code == 200
        assert r.json()["updated"] == 0


# --- Excel export ---
class TestExcelExport:
    def test_export_returns_xlsx(self, session):
        # Ensure at least one record exists
        c = session.post(f"{API}/workorders", json={"ot_number": "TEST_EXCEL_001", "requestor": "Excel User", "task_detail": "Detail"}).json()
        try:
            r = session.get(f"{API}/workorders/export/excel")
            assert r.status_code == 200
            ct = r.headers.get("content-type", "")
            assert "spreadsheetml.sheet" in ct or "xlsx" in ct.lower()
            # Parse xlsx
            wb = load_workbook(io.BytesIO(r.content))
            assert "Órdenes de Trabajo" in wb.sheetnames
            ws = wb["Órdenes de Trabajo"]
            headers = [cell.value for cell in ws[1]]
            expected_headers = [
                'Número de OT', 'Fecha y Hora', 'Solicitante',
                'Detalle de la Tarea', 'Número de Service Desk',
                'Observaciones', 'Tiene Adjunto', 'Nombre del Archivo'
            ]
            for h in expected_headers:
                assert h in headers, f"Missing header: {h}"
        finally:
            session.delete(f"{API}/workorders/{c['id']}")


# --- Attachment upload/download ---
class TestAttachments:
    def test_upload_and_download(self, session, created_wo):
        file_content = b"hello attachment world"
        files = {"file": ("test.txt", io.BytesIO(file_content), "text/plain")}
        # multipart upload (don't use json Content-Type)
        up = requests.post(f"{API}/workorders/{created_wo['id']}/upload", files=files)
        assert up.status_code == 200, up.text
        assert up.json()["filename"] == "test.txt"

        # download
        dl = requests.get(f"{API}/workorders/{created_wo['id']}/attachment")
        assert dl.status_code == 200
        assert dl.content == file_content

    def test_download_no_attachment(self, session):
        c = session.post(f"{API}/workorders", json={"ot_number": "TEST_NOATT_001"}).json()
        try:
            r = requests.get(f"{API}/workorders/{c['id']}/attachment")
            assert r.status_code == 404
        finally:
            session.delete(f"{API}/workorders/{c['id']}")
